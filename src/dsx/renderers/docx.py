"""Renderer DOCX berbasis python-docx.

Catatan pemetaan:
- Style DSX dipetakan ke style paragraf Word. Style bawaan (Normal, Title,
  Subtitle, Heading 1-9, Quote, Caption) dimodifikasi; style lain dibuat baru.
- Properti lama pada style bawaan dibersihkan lebih dulu agar hasil sama persis
  dengan style hasil resolusi (mis. border Title, atribut tema font).
- Heading mendapat identitas struktural: style Heading N atau style custom yang
  berbasis Heading N, ditambah outline level pada paragraf.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor, Twips

from ..ir.nodes import Callout, DocumentIR, Figure, HeaderFooter, ListBlock, MathBlock, MathSeg, PageNumber, Paragraph, Table, TableCell, TextBlock, TextSeg, LinkSeg, FootnoteSeg, TocBlock
from ..limits import Deadline
from ..styles.defaults import CALLOUT_THEMES
from ..styles.registry import ResolvedStyle
from .base import Renderer
from .omml import build_inline_math, build_math_para

_FIXED_TIMESTAMP = datetime(2000, 1, 1, 0, 0, 0)
_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_TABLE_ALIGN = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}
_BUILTIN_NAMES = {
    "normal": "Normal",
    "title": "Title",
    "subtitle": "Subtitle",
    "quote": "Quote",
    "caption": "Caption",
}
for _n in range(1, 10):
    _BUILTIN_NAMES[f"heading-{_n}"] = f"Heading {_n}"
_NUMERIC = re.compile(r"^[+-]?\d+([.,]\d+)?$")
_FIELD_SPLIT = re.compile(r"(\{pages?\})")
# Teks sel tabel satu spasi (tidak mewarisi jarak baris isi dokumen, mis. 1.5).
TABLE_LINE_SPACING = 1.0
# Tabel dengan baris (termasuk header) <= nilai ini tidak dipecah antarhalaman.
TABLE_KEEP_TOGETHER_MAX_ROWS = 8

_FONT_KEYS = ("font", "size", "bold", "italic", "underline", "color")


class DocxRenderError(RuntimeError):
    pass



def _fill_segments(par: Any, segments: tuple, plain: str, math_size: float | None = None) -> list[Any]:
    """Isi paragraf dengan segmen kaya atau teks polos; kembalikan daftar run."""
    runs: list[Any] = []
    if segments:
        for seg in segments:
            if isinstance(seg, TextSeg):
                run = par.add_run(seg.text)
                if seg.bold:
                    run.bold = True
                if seg.italic:
                    run.italic = True
                if seg.underline:
                    run.underline = True
                runs.append(run)
            elif isinstance(seg, LinkSeg):
                run = _add_hyperlink(par, seg.text, seg.url)
                if seg.bold:
                    run.bold = True
                if seg.italic:
                    run.italic = True
                run.underline = True
                runs.append(run)
            elif isinstance(seg, FootnoteSeg):
                runs.append(_add_footnote_ref(par, seg.note_id))
            elif isinstance(seg, MathSeg):
                par._p.append(parse_xml(etree_tostring(build_inline_math(seg.tree, math_size))))
            else:
                # fallback
                runs.append(par.add_run(getattr(seg, "text", "")))
    elif plain:
        runs.append(par.add_run(plain))
    return runs

class DocxRenderer(Renderer):
    format = "docx"

    def render(self, ir: DocumentIR, deadline: Deadline | None = None) -> bytes:
        self._deadline = deadline
        doc = Document()
        body = doc._element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        self._style_map: dict[str, Any] = {}
        self._numbering = _Numbering(doc)
        self._set_properties(doc, ir)
        self._setup_page(doc, ir)
        self._materialize_styles(doc, ir)
        self._set_language(doc, ir)
        for block in ir.blocks:
            if self._deadline is not None:
                self._deadline.check()
            if isinstance(block, Paragraph):
                self._paragraph(doc, block)
            elif isinstance(block, TextBlock):
                for p in block.paragraphs:
                    self._paragraph(doc, p)
            elif isinstance(block, Figure):
                self._figure(doc, block)
            elif isinstance(block, Table):
                self._table(doc, block, ir)
            elif isinstance(block, ListBlock):
                self._list(doc, block)
            elif isinstance(block, Callout):
                self._callout(doc, block, ir)
            elif isinstance(block, MathBlock):
                self._math(doc, block)
            elif isinstance(block, TocBlock):
                self._toc(doc, block)
        if not ir.blocks:
            doc.add_paragraph(style=self._style_map["normal"])
        self._header_footer(doc, ir)
        self._footnotes(doc, ir)
        buffer = io.BytesIO()
        doc.save(buffer)
        return _normalize_zip(buffer.getvalue())

    # ------------------------------------------------------------- document
    @staticmethod
    def _set_properties(doc: Any, ir: DocumentIR) -> None:
        props = doc.core_properties
        props.title = ir.metadata.title
        props.author = ir.metadata.author
        props.subject = ir.metadata.subject
        props.last_modified_by = ""
        props.comments = ""
        props.keywords = ""
        props.category = ""
        props.revision = 1
        props.created = _FIXED_TIMESTAMP
        props.modified = _FIXED_TIMESTAMP

    @staticmethod
    def _setup_page(doc: Any, ir: DocumentIR) -> None:
        page = ir.page
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE if page.orientation == "landscape" else WD_ORIENT.PORTRAIT
        section.page_width = Cm(page.width_cm)
        section.page_height = Cm(page.height_cm)
        section.top_margin = Cm(page.margin_top)
        section.bottom_margin = Cm(page.margin_bottom)
        section.left_margin = Cm(page.margin_left)
        section.right_margin = Cm(page.margin_right)
        section.header_distance = Cm(min(1.25, page.margin_top / 2))
        section.footer_distance = Cm(min(1.25, page.margin_bottom / 2))
        section.gutter = Cm(0)

    @staticmethod
    def _set_language(doc: Any, ir: DocumentIR) -> None:
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        if defaults is None:
            return
        rpr_default = defaults.find(qn("w:rPrDefault"))
        rpr = rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
        if rpr is None:
            return
        normal = ir.styles["normal"].props["font"]
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in list(fonts.attrib):
            del fonts.attrib[attr]
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), normal)
        if ir.metadata.language:
            lang = rpr.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang")
                rpr.append(lang)
            for attr in list(lang.attrib):
                del lang.attrib[attr]
            lang.set(qn("w:val"), ir.metadata.language)

    # --------------------------------------------------------------- styles
    def _materialize_styles(self, doc: Any, ir: DocumentIR) -> None:
        existing = {s.name.casefold() for s in doc.styles}
        existing_ids = {s.style_id for s in doc.styles}
        order = ["normal"] + [n for n in sorted(ir.styles) if n != "normal"]
        for name in order:
            rs = ir.styles[name]
            docx_name = _BUILTIN_NAMES.get(name)
            if docx_name is not None:
                try:
                    style = doc.styles[docx_name]
                except KeyError:
                    raise DocxRenderError(f"Style bawaan Word {docx_name!r} tidak tersedia pada template.") from None
            else:
                candidate = name
                counter = 0
                while candidate.casefold() in existing or candidate.replace(" ", "") in existing_ids:
                    counter += 1
                    candidate = f"DSX {name}" if counter == 1 else f"DSX {counter} {name}"
                style = doc.styles.add_style(candidate, WD_STYLE_TYPE.PARAGRAPH)
                existing.add(candidate.casefold())
                existing_ids.add(style.style_id)
                if rs.role == "heading" and rs.level is not None:
                    style.base_style = doc.styles[f"Heading {rs.level}"]
                else:
                    style.base_style = doc.styles["Normal"]
                style.quick_style = True
            self._apply_style(style, rs)
            self._style_map[name] = style

    @staticmethod
    def _apply_style(style: Any, rs: ResolvedStyle) -> None:
        element = style.element
        for tag in ("w:pPr", "w:rPr", "w:link"):
            for child in element.findall(qn(tag)):
                element.remove(child)
        props = rs.props
        font = style.font
        font.name = props["font"]
        rfonts = element.rPr.rFonts
        rfonts.set(qn("w:eastAsia"), props["font"])
        rfonts.set(qn("w:cs"), props["font"])
        font.size = Pt(props["size"])
        font.bold = props["bold"]
        font.italic = props["italic"]
        font.underline = props["underline"]
        font.color.rgb = RGBColor.from_string(props["color"].lstrip("#").upper())
        pf = style.paragraph_format
        pf.alignment = _ALIGN[props["align"]]
        pf.line_spacing = float(props["line"])
        pf.space_before = Pt(props["space_before"])
        pf.space_after = Pt(props["space_after"])
        pf.left_indent = Cm(props["indent_left"])
        pf.right_indent = Cm(props["indent_right"])
        pf.first_line_indent = Cm(props["indent_first"])
        pf.page_break_before = bool(props["page_break_before"])
        if rs.role == "heading":
            pf.keep_with_next = True
            pf.keep_together = True
            if rs.level is not None:
                _set_outline_level(style.element.get_or_add_pPr(), rs.level)

    # ----------------------------------------------------------- paragraphs
    def _paragraph(self, container: Any, p: Paragraph) -> Any:
        par = container.add_paragraph(style=self._style_map[p.style])
        runs = _fill_segments(par, p.segments, p.text, p.overrides.get("size"))
        _apply_overrides(par, runs, p.overrides)
        if p.is_heading and p.level is not None:
            _set_outline_level(par._p.get_or_add_pPr(), p.level)
        return par

    # --------------------------------------------------------------- figure
    def _figure(self, doc: Any, fig: Figure) -> None:
        par = doc.add_paragraph(style=self._style_map["normal"])
        overrides: dict[str, Any] = {
            "align": fig.align,
            "indent_first": 0,
            "indent_left": 0,
            "indent_right": 0,
        }
        if fig.page_break_before:
            overrides["page_break_before"] = True
        run = par.add_run()
        shape = run.add_picture(io.BytesIO(fig.image.data), width=Cm(fig.width_cm), height=Cm(fig.height_cm))
        if fig.alt:
            shape._inline.docPr.set("descr", fig.alt)
        _apply_overrides(par, [], overrides)
        if fig.caption is not None:
            par.paragraph_format.keep_with_next = True
            self._paragraph(doc, fig.caption)

    # ---------------------------------------------------------------- table
    def _table(self, doc: Any, tbl: Table, ir: DocumentIR) -> None:
        if tbl.page_break_before and tbl.caption is None:
            # page break before table: empty paragraph with break (cell pPr break is unreliable)
            br = doc.add_paragraph(style=self._style_map["normal"])
            br.paragraph_format.page_break_before = True
            br.paragraph_format.space_before = Pt(0)
            br.paragraph_format.space_after = Pt(0)
        if tbl.caption is not None:
            cap = self._paragraph(doc, tbl.caption)
            cap.paragraph_format.keep_with_next = True
            if tbl.page_break_before:
                cap.paragraph_format.page_break_before = True
        ncols = tbl.column_count
        table = doc.add_table(rows=tbl.row_count, cols=ncols)
        table.style = doc.styles["Table Grid"]
        table.alignment = _TABLE_ALIGN[tbl.align]
        table.autofit = False
        col_cm = ir.page.content_width / ncols
        col_twips = int(round(Cm(col_cm).twips))
        col_width = Cm(col_cm)
        # Fixed layout + explicit grid so columns stay equal in Word/LibreOffice/PDF
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(col_twips * ncols))
        # Rebuild tblGrid
        tbl_el = table._tbl
        existing_grid = tbl_el.find(qn("w:tblGrid"))
        if existing_grid is not None:
            tbl_el.remove(existing_grid)
        grid = OxmlElement("w:tblGrid")
        for _ in range(ncols):
            grid_col = OxmlElement("w:gridCol")
            grid_col.set(qn("w:w"), str(col_twips))
            grid.append(grid_col)
        # insert grid after tblPr
        tbl_pr.addnext(grid)
        for column in table.columns:
            column.width = col_width
        all_rows = ([tbl.header] if tbl.header is not None else []) + list(tbl.rows)
        normal = self._style_map["normal"]
        keep_together = len(all_rows) <= TABLE_KEEP_TOGETHER_MAX_ROWS
        for r_index, values in enumerate(all_rows):
            if self._deadline is not None and r_index % 200 == 0:
                self._deadline.check()
            is_header = tbl.header is not None and r_index == 0
            row = table.rows[r_index]
            tr_pr = row._tr.get_or_add_trPr()
            # avoid duplicate cantSplit/tblHeader if re-run
            for tag in ("w:cantSplit", "w:tblHeader"):
                for old in tr_pr.findall(qn(tag)):
                    tr_pr.remove(old)
            tr_pr.append(OxmlElement("w:cantSplit"))
            if is_header:
                tr_pr.append(OxmlElement("w:tblHeader"))
            # pad short rows (defensive; validator normally rejects mismatch)
            if len(values) < ncols:
                values = list(values) + [""] * (ncols - len(values))
            for c_index in range(ncols):
                raw = values[c_index] if c_index < len(values) else ""
                if isinstance(raw, TableCell):
                    cell_obj = raw
                else:
                    cell_obj = TableCell(str(raw) if raw is not None else "", ())
                cell = row.cells[c_index]
                cell.width = col_width
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.insert(0, tc_w)
                tc_w.set(qn("w:type"), "dxa")
                tc_w.set(qn("w:w"), str(col_twips))
                v_align = tc_pr.find(qn("w:vAlign"))
                if v_align is None:
                    v_align = OxmlElement("w:vAlign")
                    tc_pr.append(v_align)
                v_align.set(qn("w:val"), "center")
                if is_header:
                    shd = tc_pr.find(qn("w:shd"))
                    if shd is None:
                        shd = OxmlElement("w:shd")
                        tc_pr.append(shd)
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D9D9D9")
                par = cell.paragraphs[0]
                par.style = normal
                runs = _fill_segments(par, cell_obj.segments, cell_obj.text)
                if is_header:
                    for run in runs:
                        try:
                            run.bold = True
                        except Exception:
                            pass
                is_numeric = (not is_header) and bool(_NUMERIC.match(cell_obj.text.strip()))
                overrides: dict[str, Any] = {
                    "align": "center" if is_header else ("right" if is_numeric else "left"),
                    "indent_first": 0,
                    "indent_left": 0,
                    "indent_right": 0,
                    "space_before": 0,
                    "space_after": 0,
                    "line": TABLE_LINE_SPACING,
                }
                _apply_overrides(par, runs, overrides)
                if keep_together and r_index < len(all_rows) - 1:
                    par.paragraph_format.keep_with_next = True
        self._spacer(doc)

    # ----------------------------------------------------------------- lists
    def _list(self, doc: Any, block: ListBlock) -> None:
        kind = "ordered" if block.ordered else "bullet"
        step, hanging = int(round(Cm(block.indent_cm).twips)), int(round(Cm(block.hanging_cm).twips))
        num_id = self._numbering.num_for(kind, step, hanging, block.start if block.ordered else None)
        style = self._style_map[block.style]
        first_paragraph = True
        for item in block.items:
            left = step * (item.level + 1)
            par = doc.add_paragraph(style=style)
            _fill_segments(par, item.segments, item.text)
            num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = item.level
            num_pr.get_or_add_numId().val = num_id
            par.paragraph_format.left_indent = Twips(left)
            par.paragraph_format.first_line_indent = Twips(-hanging)
            par.paragraph_format.keep_together = True  # item tidak terpotong antarhalaman
            par.paragraph_format.keep_with_next = bool(item.continuations)
            overrides = dict(block.overrides)
            if block.page_break_before and first_paragraph:
                overrides["page_break_before"] = True
            first_paragraph = False
            _apply_overrides(par, par.runs, overrides)
            for index, extra in enumerate(item.continuations):
                cont = doc.add_paragraph(style=style)
                cont.add_run(extra)
                cont.paragraph_format.left_indent = Twips(left)  # sejajar dengan teks item, tanpa nomor
                cont.paragraph_format.first_line_indent = Twips(0)
                cont.paragraph_format.keep_together = True
                cont.paragraph_format.keep_with_next = index < len(item.continuations) - 1
                _apply_overrides(cont, cont.runs, {k: v for k, v in block.overrides.items() if k not in ("indent_first",)})

    # ---------------------------------------------------------------- callout
    def _callout(self, doc: Any, block: Callout, ir: DocumentIR) -> None:
        theme = CALLOUT_THEMES[block.variant]
        paragraphs = ([block.title] if block.title is not None else []) + list(block.body)
        table = doc.add_table(rows=1, cols=1)
        table.style = doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        width = Cm(ir.page.content_width)
        table.columns[0].width = width
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is not None:
            tbl_w.set(qn("w:type"), "dxa")
            tbl_w.set(qn("w:w"), str(int(round(width.twips))))
        label = block.title.text if block.title is not None else block.variant
        for tag, value in (("w:tblCaption", f"dsx-callout:{block.variant}"), ("w:tblDescription", label)):
            element = OxmlElement(tag)
            element.set(qn("w:val"), value)
            tbl_pr.append(element)
        row = table.rows[0]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cell = row.cells[0]
        cell.width = width
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{theme["border"]}"/>'
            f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{theme["border"]}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{theme["border"]}"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{theme["border"]}"/>'
            "</w:tcBorders>"
        )
        tc_pr.append(borders)
        tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{theme["fill"]}"/>'))
        tc_pr.append(
            parse_xml(
                f'<w:tcMar {nsdecls("w")}><w:top w:w="100" w:type="dxa"/><w:left w:w="180" w:type="dxa"/>'
                '<w:bottom w:w="100" w:type="dxa"/><w:right w:w="180" w:type="dxa"/></w:tcMar>'
            )
        )
        for index, p in enumerate(paragraphs):
            par = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            par.style = self._style_map[p.style]
            runs = _fill_segments(par, p.segments, p.text, p.overrides.get("size"))
            overrides = dict(p.overrides)
            if block.page_break_before and index == 0:
                overrides["page_break_before"] = True
            _apply_overrides(par, runs, overrides)
            if block.title is not None and index == 0 and len(paragraphs) > 1:
                par.paragraph_format.keep_with_next = True
        self._spacer(doc)

    # ------------------------------------------------------------------- math
    def _math(self, doc: Any, block: MathBlock) -> None:
        par = doc.add_paragraph(style=self._style_map[block.style])
        overrides = dict(block.overrides)
        if block.page_break_before:
            overrides["page_break_before"] = True
        _apply_overrides(par, [], overrides)
        par.paragraph_format.keep_together = True
        omml = build_math_para(block.tree, block.align, block.size)
        par._p.append(parse_xml(etree_tostring(omml)))

    def _spacer(self, doc: Any) -> None:
        spacer = doc.add_paragraph(style=self._style_map["normal"])
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.first_line_indent = Cm(0)
        spacer.paragraph_format.line_spacing = Pt(8)

    def _toc(self, doc: Any, block: TocBlock) -> None:
        if block.title:
            title = doc.add_paragraph(block.title)
            title.style = self._style_map.get("heading-1", self._style_map["normal"])
            if block.page_break_before:
                title.paragraph_format.page_break_before = True
        par = doc.add_paragraph()
        run = par.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        run._r.append(fld_begin)
        run2 = par.add_run()
        instr = OxmlElement("w:instrText")
        instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        instr.text = f' TOC \\o "1-{block.depth}" \\h \\z \\u '
        run2._r.append(instr)
        run3 = par.add_run()
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        run3._r.append(fld_sep)
        par.add_run("(Daftar isi — perbarui field di Word/LibreOffice)")
        run4 = par.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run4._r.append(fld_end)
        self._spacer(doc)

    def _footnotes(self, doc: Any, ir: DocumentIR) -> None:
        if not ir.footnotes:
            return
        sep = doc.add_paragraph()
        sep_run = sep.add_run("─" * 24)
        sep_run.font.size = Pt(10)
        heading = doc.add_paragraph()
        hr = heading.add_run("Catatan kaki")
        hr.bold = True
        hr.font.size = Pt(11)
        for fn in ir.footnotes:
            par = doc.add_paragraph()
            mark = par.add_run(f"[{fn.note_id}] ")
            mark.bold = True
            mark.font.size = Pt(9)
            body = par.add_run(fn.text)
            body.font.size = Pt(9)

    # ------------------------------------------------------ header / footer
    def _header_footer(self, doc: Any, ir: DocumentIR) -> None:
        section = doc.sections[0]
        header_items: list[Any] = []
        footer_items: list[Any] = []
        if ir.header is not None:
            header_items.append(ir.header)
        if ir.footer is not None:
            footer_items.append(ir.footer)
        if ir.page_number is not None:
            (header_items if ir.page_number.position == "header" else footer_items).append(ir.page_number)
        for part, items in ((section.header, header_items), (section.footer, footer_items)):
            if not items:
                continue
            part.is_linked_to_previous = False
            first = True
            for item in items:
                par = part.paragraphs[0] if first else part.add_paragraph()
                first = False
                par.style = self._style_map[item.style]
                runs: list[Any] = []
                if isinstance(item, HeaderFooter):
                    runs.append(par.add_run(item.text))
                else:
                    runs.extend(_page_number_runs(par, item.format))
                _apply_overrides(par, runs, item.overrides)


# ---------------------------------------------------------------------- util
def _set_outline_level(ppr: Any, level: int) -> None:
    for existing in ppr.findall(qn("w:outlineLvl")):
        ppr.remove(existing)
    element = OxmlElement("w:outlineLvl")
    element.set(qn("w:val"), str(level - 1))
    ppr.insert_element_before(element, "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _apply_overrides(par: Any, runs: list[Any], overrides: dict[str, Any]) -> None:
    pf = par.paragraph_format
    for key, value in overrides.items():
        if key == "align":
            pf.alignment = _ALIGN[value]
        elif key == "line":
            pf.line_spacing = float(value)
        elif key == "space_before":
            pf.space_before = Pt(value)
        elif key == "space_after":
            pf.space_after = Pt(value)
        elif key == "indent_left":
            pf.left_indent = Cm(value)
        elif key == "indent_right":
            pf.right_indent = Cm(value)
        elif key == "indent_first":
            pf.first_line_indent = Cm(value)
        elif key == "page_break_before":
            pf.page_break_before = bool(value)
    if not any(k in overrides for k in _FONT_KEYS):
        return
    for run in runs:
        font = run.font
        if "font" in overrides:
            font.name = overrides["font"]
            rfonts = run._r.rPr.rFonts
            rfonts.set(qn("w:eastAsia"), overrides["font"])
            rfonts.set(qn("w:cs"), overrides["font"])
        if "size" in overrides:
            font.size = Pt(overrides["size"])
        if "bold" in overrides:
            font.bold = overrides["bold"]
        if "italic" in overrides:
            font.italic = overrides["italic"]
        if "underline" in overrides:
            font.underline = overrides["underline"]
        if "color" in overrides:
            font.color.rgb = RGBColor.from_string(overrides["color"].lstrip("#").upper())


def _field_run(par: Any, kind: str, text: str | None = None) -> Any:
    run = par.add_run()
    if kind in ("begin", "separate", "end"):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        run._r.append(element)
    elif kind == "instr":
        element = OxmlElement("w:instrText")
        element.set(qn("xml:space"), "preserve")
        element.text = text
        run._r.append(element)
    return run


def _page_number_runs(par: Any, fmt: str) -> list[Any]:
    runs: list[Any] = []
    for part in _FIELD_SPLIT.split(fmt):
        if not part:
            continue
        if part in ("{page}", "{pages}"):
            instr = "PAGE" if part == "{page}" else "NUMPAGES"
            runs.append(_field_run(par, "begin"))
            runs.append(_field_run(par, "instr", f" {instr} "))
            runs.append(_field_run(par, "separate"))
            runs.append(par.add_run("1"))
            runs.append(_field_run(par, "end"))
        else:
            runs.append(par.add_run(part))
    return runs


def _normalize_zip(data: bytes) -> bytes:
    """Tulis ulang ZIP dengan timestamp tetap agar keluaran deterministik."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as dst:
        for info in src.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.create_system = 3
            entry.external_attr = 0o600 << 16
            dst.writestr(entry, src.read(info))
    return out.getvalue()


def etree_tostring(element: Any) -> bytes:
    from lxml import etree

    return etree.tostring(element, encoding="utf-8")


_BULLETS = ("\u2022", "\u25e6", "\u25aa")
_ORDERED_FORMATS = ("decimal", "lowerLetter", "lowerRoman")


class _Numbering:
    """Definisi penomoran native Word (numbering.xml).

    Satu abstractNum per kombinasi (jenis, indentasi, hanging); bullet berbagi satu w:num,
    sedangkan setiap daftar berurutan mendapat w:num sendiri dengan startOverride agar
    penomoran selalu mulai ulang. Memakai XML internal karena python-docx tidak memiliki
    API publik untuk numbering.
    """

    def __init__(self, doc: Any) -> None:
        self._root = doc.part.numbering_part.element
        self._abstracts: dict[tuple, int] = {}
        self._bullet_nums: dict[tuple, int] = {}

    def _next_id(self, tag: str, attr: str) -> int:
        ids = [int(e.get(qn(attr))) for e in self._root.findall(qn(tag))]
        return (max(ids) + 1) if ids else 1

    def _add_abstract(self, kind: str, step: int, hanging: int) -> int:
        abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        levels = []
        for i in range(9):
            left = step * (i + 1)
            if kind == "bullet":
                fmt, text = "bullet", _BULLETS[i % 3]
                rpr = '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri" w:hint="default"/></w:rPr>'
            else:
                fmt, text, rpr = _ORDERED_FORMATS[i % 3], f"%{i + 1}.", ""
            levels.append(
                f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>{rpr}</w:lvl>'
            )
        element = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            '<w:multiLevelType w:val="hybridMultilevel"/>' + "".join(levels) + "</w:abstractNum>"
        )
        first_num = self._root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(element)
        else:
            self._root.append(element)
        return abstract_id

    def _add_num(self, abstract_id: int, start: int | None = None) -> int:
        num_id = self._next_id("w:num", "w:numId")
        override = (
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{start}"/></w:lvlOverride>' if start is not None else ""
        )
        self._root.append(
            parse_xml(f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>')
        )
        return num_id

    def num_for(self, kind: str, step: int, hanging: int, start: int | None) -> int:
        key = (kind, step, hanging)
        if key not in self._abstracts:
            self._abstracts[key] = self._add_abstract(kind, step, hanging)
        if kind == "bullet":
            if key not in self._bullet_nums:
                self._bullet_nums[key] = self._add_num(self._abstracts[key])
            return self._bullet_nums[key]
        return self._add_num(self._abstracts[key], start)


def _add_hyperlink(paragraph: Any, text: str, url: str) -> Any:
    """Tambah hyperlink eksternal; mengembalikan proxy run untuk bold/italic."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    class _RunProxy:
        def __init__(self, element: Any) -> None:
            self._r = element

        @property
        def bold(self) -> bool | None:
            return None

        @bold.setter
        def bold(self, value: bool) -> None:
            rPr = self._r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                self._r.insert(0, rPr)
            existing = rPr.find(qn("w:b"))
            if value and existing is None:
                rPr.append(OxmlElement("w:b"))
            elif not value and existing is not None:
                rPr.remove(existing)

        @property
        def italic(self) -> bool | None:
            return None

        @italic.setter
        def italic(self, value: bool) -> None:
            rPr = self._r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                self._r.insert(0, rPr)
            existing = rPr.find(qn("w:i"))
            if value and existing is None:
                rPr.append(OxmlElement("w:i"))
            elif not value and existing is not None:
                rPr.remove(existing)

        @property
        def underline(self) -> bool | None:
            return True

        @underline.setter
        def underline(self, value: bool) -> None:
            pass

        @property
        def font(self) -> Any:
            class F:
                name = None
                size = None
                color = type("C", (), {"rgb": None})()
            return F()

    return _RunProxy(new_run)


def _add_footnote_ref(paragraph: Any, note_id: int) -> Any:
    run = paragraph.add_run(f"[{note_id}]")
    run.font.superscript = True
    run.font.size = Pt(9)
    return run
